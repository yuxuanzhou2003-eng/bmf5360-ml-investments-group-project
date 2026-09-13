from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"D:\3 Study 学习资料\2E 金融研二(上)资料\BMF5360 Machine Learning in Investments\Group project_2.0")
OUT = ROOT / "deliverables" / "BMF5360_Baseline_Package" / "BMF5360_AI_Baseline_Model_Report.docx"
NAVY, PALE, BORDER = "17365D", "F5F7F9", "D9D9D9"

def font(run, size=10.5, bold=False, color=None):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def cell_text(cell, value, size=8.5, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    font(p.add_run(str(value)), size, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, headers, rows, widths, size=8.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        shade(t.rows[0].cells[i], NAVY)
        cell_text(t.rows[0].cells[i], h, size, True, "FFFFFF")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            if ri % 2: shade(cells[i], PALE)
            cell_text(cells[i], val, size)
    for row in t.rows:
        for i, width in enumerate(widths): row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.24
    font(p.add_run(text))

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(11 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    font(p.add_run(text), 15 if level == 1 else 12)

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(.72)
sec.left_margin = sec.right_margin = Inches(.78)
for name, size in [("Normal", 10.5), ("Title", 22), ("Heading 1", 15), ("Heading 2", 12)]:
    st = doc.styles[name]
    st.font.name = "Arial"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(footer.add_run("BMF5360 | AI supply-chain daily baseline | development and validation evidence only"), 8)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("BMF5360 AI产业链日频基准模型说明"), 22)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("机器学习投资课程项目 | 可复核的开发期与验证期材料")
font(r, 11); r.italic = True
paragraph(doc, "本文件说明当前基准策略的研究问题、公司池、因子、标签、时间切分、模型和验证结果。它是开发期和验证期的研究记录，不是基金业绩承诺。2023-01-01 至 2026-06-30 的最终测试标签仍然封存，未在本报告中打开或用于模型选择。")
table(doc, ["项目", "当前定义"], [
    ["研究问题", "在既定的美国上市 AI 产业链股票池内，预测未来 21 个交易日是否跑赢 SPY。"],
    ["建模基准", "技术 Logistic 与技术加宏观 Logistic 的预先定义分类比较。"],
    ["开发和验证", "训练 2015-2020；验证 2021-2022；最终测试期封存。"],
    ["核心验证结果", "Pooled ROC-AUC 0.5616；但择时与选股分解后，宏观的横截面增量未获验证。"],
], [1.45, 5.7], 9)

heading(doc, "1 研究问题和基金定位")
paragraph(doc, "策略的范围是美国上市 AI 产业链公司，而不是把所有使用 AI 的企业都视为可投资对象。研究目标是检验：截止形成日可获得的技术、流动性、市场和宏观状态信息，能否帮助识别未来阶段相对 SPY 更可能获胜的股票。")
paragraph(doc, "当前模型输出的是分类分数。只有在后续固定交易规则、成本假设和组合约束后，分数才会转化为持仓；因此 AUC 不是收益率、夏普或已验证的基金业绩。")

heading(doc, "2 公司选择逻辑")
paragraph(doc, "公司池先按经济角色定义，再对照 781 只历史 S&P 500 RIC 的母表。筛选没有使用未来收益、未来估值或模型表现。每家公司只分配一个主组；次要角色仅作记录，避免同一公司在组内重复计算。退市不是剔除条件：JNPR.N^G25 及其价格历史保留在登记表中，是否可在某日交易由成员资格和当日数据可用性决定。")
table(doc, ["主组", "经济角色", "公司", "直接 PIT", "暂定静态"], [
    ["gpu_accelerator", "GPU、CPU、定制加速器", "5", "2", "3"],
    ["ai_semiconductor", "晶圆、EDA、功率/连接芯片及设备", "6", "0", "6"],
    ["memory_hbm_storage", "HBM、DRAM、NAND 与企业存储", "5", "1", "4"],
    ["server_network", "AI 服务器、交换机、光互联", "6", "1", "5"],
    ["cloud_software", "云平台、数据库、企业软件和开发工具", "10", "1", "9"],
    ["data_center_power_cooling", "数据中心、供电、制冷和基础设施", "9", "3", "6"],
    ["robotics_autonomy", "工业自动化、机器视觉和自动驾驶", "8", "1", "7"],
    ["合计", "预先定义的 AI 产业链探索池", "49", "9", "40"],
], [1.35, 3.25, .55, .85, 1.0], 8)
paragraph(doc, "证据质量必须与经济分类分开表述。9 家公司有本地、带日期的披露证据；其余 40 家只有当前官方描述的暂定静态证据，不能被包装为历史时点的完整证明。报告因此将 49 家称为探索池，并把建设严格历史可投资池列为下一步。")

heading(doc, "3 数据处理和因子选择")
table(doc, ["数据层", "处理规则", "影响"], [
    ["LSEG 日频价格与报价", "保留供应商已调整价格/收益，不做二次拆股调整。", "缺失保留为 NA，原始路径与覆盖情况可审计。"],
    ["技术与市场变量", "使用形成日 F-1 的滚动窗口；历史不足时保留缺失。", "不前填、不以零替代缺失。"],
    ["FRED 宏观变量", "按 SPY 日历对齐，在 F 使用 F-1 可得信息。", "保留来源日期；完整发布 vintage 是限制。"],
    ["信用利差候选", "训练/验证窗口无有效起始覆盖，6 个变量全数排除。", "不以未来数据填补；首个有效日期为 2023-09-11。"],
], [1.3, 3.25, 2.3], 8.5)
paragraph(doc, "因子按投资解释预先分组，而不是按验证 AUC 逐个筛选。21 个技术、流动性和市场变量刻画动量、风险、成交活跃度、报价成本、beta 和 SPY 状态；20 个宏观变量刻画风险偏好、利率曲线和美元环境。每个变量的公式、单位、窗口和截止时点都在 Excel 的 Factor Dictionary 工作表逐项列示。")
table(doc, ["类别", "数量", "选择逻辑"], [
    ["个股动量", "4", "1、5、20、60 个 SPY 交易日复合收益。"],
    ["个股风险与流动性", "11", "波动、成交量/成交额、买卖价差、beta、特质波动与交易天数。"],
    ["市场状态", "6", "SPY 动量和波动，描述共同市场环境。"],
    ["宏观状态", "20", "VIX、短中长利率、美元和期限利差。"],
    ["已排除信用利差", "6", "训练和验证期未覆盖，避免以未来数据填补。"],
], [1.55, .65, 4.65], 8.5)
paragraph(doc, "模型管线中的缺失处理只在训练集拟合：中位数填补加缺失指示变量。它不会改变原始数据，也不会把 NA 当作零。宏观变量在没有交互项时不能直接改变同日个股的线性排序；若要研究公司对宏观状态的差异反应，应预先限定少量交互项，或将总体仓位择时与个股选股分成两个模块。")

heading(doc, "4 标签 时间切分和泄露防护")
paragraph(doc, "目标变量 y=1 的条件是：从形成日后的未来 21 个 SPY 交易日，个股总回报减去 SPY 总回报大于零。所有特征仅使用形成日 F 之前可见的信息；未来收益只用于在持有期结束后生成标签。")
table(doc, ["阶段", "期间", "锚点", "公司", "规则"], [
    ["训练", "2015-2020", "2,334", "39", "按公司采用 H21 非重叠锚点。"],
    ["验证", "2021-2022", "992", "44", "同一验证样本比较冻结的基准模型。"],
    ["最终测试", "2023-01 至 2026-06", "未打开", "未打开", "不读标签、不生成预测或指标。"],
], [1.0, 1.45, .85, .65, 3.0], 8.5)
paragraph(doc, "训练内部的随机森林调参使用扩展窗口，并要求训练标签退出日期早于内部验证起点。非重叠锚点降低同一公司持有期重叠，但不同股票仍可能共享市场冲击；后续置信区间应按形成日时间块处理，不能把 992 条记录当作独立观察。")

heading(doc, "5 基准模型和验证结果")
paragraph(doc, "Logistic 回归是监督式机器学习分类器，也是本项目应保留的主要基准：它参数少、方向可审阅，并能检验复杂模型是否真的带来增量。Random Forest 已作为非线性对照，而不是被假定为必然更好。管线为 SimpleImputer median add_indicator keep_empty_features，随后 StandardScaler，再以 C=1、class_weight=balanced、lbfgs、max_iter=2000 和 random_state=5360 的 LogisticRegression 拟合。")
table(doc, ["特征组 / 模型", "ROC-AUC", "PR-AUC", "Brier", "准确率", "形成日高低分差"], [
    ["技术 / Logistic", "0.5395", "0.5421", "0.2511", "51.61%", "0.6152%"],
    ["技术 / Random Forest", "0.5231", "0.5511", "0.2543", "51.92%", "0.1615%"],
    ["技术加宏观 / Logistic", "0.5616", "0.5524", "0.2508", "56.45%", "0.5244%"],
    ["技术加宏观 / Random Forest", "0.5362", "0.5545", "0.2496", "52.92%", "-0.1076%"],
], [2.1, .85, .85, .75, .9, 1.1], 8.25)
paragraph(doc, "技术加宏观 Logistic 的 pooled 验证 ROC-AUC 为 0.561559，相比技术 Logistic 高 0.022096。择时与选股审计显示，41 个原始变量中 27 个在同一形成日内为常数，因而 pooled AUC 混合了共同日期状态与股票排序。按 23 个 21-session 区块重抽样，pooled AUC 差的 95% 区间为 -0.03944 至 +0.07037，不能作为宏观变量带来稳健增量的推断性证据。")
paragraph(doc, "严格在同一形成日比较股票时，技术加宏观相对技术模型的 pair-weighted AUC 增量为 0.00624，区间为 -0.00658 至 +0.02184；日内 rank-IC 增量为 0.00821，区间为 -0.00739 至 +0.02403。因此当前没有经验证的宏观选股提升。0.5616 是混合择时/选股的探索性 pooled 指标，不是 56.16% 的回报、胜率或基金能力。")
paragraph(doc, "复杂度没有自动带来价值：同一特征集的 Random Forest AUC 为 0.5362，低于 Logistic。`spread_median_60_bps` 的缺失指示变量虽在训练系数中靠前，但在这 992 条验证预测中原始价差均非缺失，因而它不驱动验证期横截面排序。当前结果支持保留 Logistic 作为透明基准，不支持已找到可稳定赚钱的复杂 ML 模型。")

heading(doc, "6 策略转化和下一步")
paragraph(doc, "下一步不是继续在验证集上搜索更多算法，而是在验证期冻结交易规则：固定每月同一交易日调仓，对当日全部合格股票打分，前 20% 等权持有到下一调仓日；同时定义单股上限、股票不足、停牌/退市、现金和成交价格的处理。")
paragraph(doc, "组合将与 SPY、同一合格池等权、简单动量、技术 Logistic 和技术加宏观 Logistic 使用相同日期范围和执行口径比较。成本按权重变动的交易金额计入，并在单边 0、5、10 和 25 bps 下做敏感性分析。验证完成、错误修正和规则冻结后，才打开一次最终测试标签。")

heading(doc, "7 材料和局限")
paragraph(doc, "本交付包包含 Excel 计算与审计表、可运行的最小代码包、模型说明和日志。Excel 的 Summary、Company Universe、Selection Logic、Factor Dictionary、Model Results、Logistic Coefficients、Processing Audit 和 Sources 工作表对应本报告的关键结论。代码包只运行四个已审计的技术/宏观基准模型，不读取封存测试标签。")
paragraph(doc, "主要局限为：49 家探索池中 40 家仍缺严格历史时点行业证据；宏观序列未使用完整发布 vintage；验证样本有共同市场冲击；当前高低分差不是基金净值；概率尚未校准；最终测试期尚未评估。材料性 AI 协助记录在 AI_USE_LOG.md；数据处理阶段、缺失处理、排除和审计状态记录在 DATA_PROCESSING_LOG.md。")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
