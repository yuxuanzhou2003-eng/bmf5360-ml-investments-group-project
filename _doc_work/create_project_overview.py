from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\3 Study 学习资料\2E 金融研二(上)资料\BMF5360 Machine Learning in Investments\Group project_2.0")
OUT = ROOT / "deliverables" / "BMF5360_项目说明_当前策略与结果_2026-09-09.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)


def set_east_asia(run, font_name="Microsoft YaHei"):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "start" if edge == "left" else "end" if edge == "right" else edge
        el = borders.find(qn(f"w:{tag}"))
        if el is None:
            el = OxmlElement(f"w:{tag}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    keep_with_next(p)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="Body Text")
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    return doc.add_paragraph(text, style=style)


def add_number(doc, text):
    return doc.add_paragraph(text, style="List Number")


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            cell.width = widths[i]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(h))
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = widths[i]
            if ridx % 2 == 1:
                set_cell_shading(cell, "F2F6FA")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(str(value))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.35)
sec.right_margin = Cm(2.35)
sec.header_distance = Cm(0.9)
sec.footer_distance = Cm(0.9)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.35
normal.paragraph_format.space_after = Pt(6)

for name in ("Body Text", "List Bullet", "List Bullet 2", "List Number"):
    st = styles[name]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(10.5)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.line_spacing = 1.35
    st.paragraph_format.space_after = Pt(5)

title = styles["Title"]
title.font.name = "Microsoft YaHei"
title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
title.font.size = Pt(25)
title.font.bold = True
title.font.color.rgb = RGBColor(0, 0, 0)

for name, size in (("Heading 1", 17), ("Heading 2", 13)):
    st = styles[name]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run("BMF5360 机器学习投资项目说明")
r.font.size = Pt(8)
r.font.color.rgb = RGBColor(100, 100, 100)
set_east_asia(r)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("项目状态截至 2026 年 9 月 9 日  |  ")
r.font.size = Pt(8)
r.font.color.rgb = RGBColor(100, 100, 100)
set_east_asia(r)
fld_char1 = OxmlElement("w:fldChar")
fld_char1.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText")
instr.set(qn("xml:space"), "preserve")
instr.text = " PAGE "
fld_char2 = OxmlElement("w:fldChar")
fld_char2.set(qn("w:fldCharType"), "end")
r._r.append(fld_char1)
r._r.append(instr)
r._r.append(fld_char2)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.paragraph_format.space_before = Pt(38)
p.paragraph_format.space_after = Pt(12)
p.add_run("从一家公司财报推断其他股票的短期反应")

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(22)
r = p.add_run("BMF5360 机器学习投资基金项目当前策略与阶段结果")
r.font.size = Pt(14)
r.bold = True
r.font.color.rgb = RGBColor(65, 65, 65)
set_east_asia(r)

add_body(doc, "这份说明写给第一次接触本项目的人。读完后，你应当能回答四个问题：我们准备交易什么、为什么一家公司发布财报会影响另一家公司、目前数据和结果到底做到哪一步、以及为什么现在还不能把这些数字当成一只已经证明有效的基金。")
add_body(doc, "目前最准确的结论是：数据采集、时间规则和两套公司关系网络已经可以审计；最初的历史价格相关性网络表现很弱；新的共同分析师网络出现了方向一致但很小的验证信号。项目值得继续推进到正式机器学习模型和组合回测，但尚未产生可向投资者承诺的收益结果。", bold_lead="目前最准确的结论是：")

add_heading(doc, "一页读懂项目", 1)
add_table(doc,
          ["问题", "现在的回答"],
          [
              ("基金交易什么", "以美国大型公司的个股为主。SPY 用作市场基准，不是当前策略的主要交易对象。"),
              ("核心想法", "一家公司公布超预期或低于预期的财报后，与它共享大量分析师覆盖的其他公司，可能在接下来的几天出现可预测的相对反应。"),
              ("机器学习做什么", "把财报意外、公司关系强度、近期价格、波动、流动性和市场环境合在一起，给候选股票排序。"),
              ("当前做到哪", "数据与关系网络已建立并通过独立检查；旧网络已完成基线模型；新网络已完成原始信号快速检验。"),
              ("当前结果", "旧网络没有带来稳定增量；新网络的验证期事件内排序相关约为 0.01，方向一致但非常弱。"),
              ("还缺什么", "新网络的正式模型、交易组合、成本与风险控制、以及最后一次密封测试。"),
          ],
          widths=[Cm(4.0), Cm(11.4)])

add_heading(doc, "我们到底在做什么", 1)
add_body(doc, "假设公司 A 今天公布业绩，利润远高于市场原先预期。这个信息首先属于 A，但它也可能告诉市场一些更广泛的事情，例如某种产品需求正在上升、某类客户的预算增加、行业成本下降，或供应链出现变化。与 A 处在同一信息圈的公司 B、C、D，价格可能不会在同一秒完成全部调整。我们的基金试图利用这段很短的信息传递过程。")
add_body(doc, "项目的预测单位不是“明天市场涨不涨”，而是“在同一次财报事件之后，哪些相关股票在接下来五个交易日相对 SPY 表现更好”。这样做让模型必须完成一个投资者真正能执行的任务：在同一批候选股票中排序，而不是只解释市场整体方向。")

add_heading(doc, "股票还是 ETF", 2)
add_body(doc, "当前核心策略是个股事件驱动。每次源公司公布财报，我们寻找最多五只相关的接收公司，并预测它们未来五个交易日相对 SPY 的收益。SPY 主要承担两个角色：一是把市场共同涨跌从目标中扣除，二是未来可用于构造市场中性的对冲。ETF 轮动仍可作为备用方案或风险覆盖层，但不是当前研究主线。")

add_heading(doc, "B C D 是怎样筛出来的", 2)
add_body(doc, "最初版本用历史股价的剩余收益相关性连接公司。这个办法直观，但它很容易把共同受大盘、行业或风险偏好影响的公司连在一起，未必代表真实的信息渠道。基线结果也显示，这类关系没有给模型带来可靠增量。")
add_body(doc, "现在的主线改为共同分析师网络。它利用一个现实机制：同一家券商或分析团队同时覆盖两家公司时，更可能持续比较它们的订单、利润率、估值和管理层指引。因此，共同覆盖可以看作一条潜在的信息与注意力通道。")
add_number(doc, "在每个财报事件之前，选择最近且严格早于公告日的分析师覆盖快照。快照距离公告最多 120 天。")
add_number(doc, "只考虑公告当天仍属于研究股票池的源公司和候选接收公司，避免事后把已经不在指数中的股票塞回历史样本。")
add_number(doc, "计算两家公司共同券商的数量和覆盖重合程度。至少需要三家共同券商。")
add_number(doc, "按覆盖重合度排序，每个财报事件最多保留五只接收公司，同时记录推荐意见重合度和仅使用可识别券商名称的重合度。")
add_number(doc, "把源公司的标准化财报意外与关系强度相乘，得到最简单的传播信号，再交给正式机器学习模型与其他历史特征共同判断。")

add_heading(doc, "一笔候选交易的时间线", 2)
add_table(doc,
          ["时间", "允许使用的信息", "目的"],
          [
              ("公告以前", "历史价格、成交量、买卖价差、历史覆盖快照和当时可见的盈利预测", "建立关系与市场背景"),
              ("公告发生", "实际盈利与公告前一致预期之间的差异", "计算财报意外"),
              ("公告后的第一个 SPY 交易日", "确定统一的正式入场会话；特征仍以入场前已知信息为准", "避免不同股票日历造成时间错位"),
              ("其后五个交易日", "只用于形成研究目标，计算接收股票相对 SPY 的累计回报", "评估预测是否有用"),
          ], widths=[Cm(3.5), Cm(8.3), Cm(3.6)])
add_body(doc, "正式的开盘或收盘执行细节、滑点和下单规则仍需在组合回测阶段冻结。目前结果只评估预测问题，没有假装已经完成真实交易。", bold_lead="尚未冻结的部分：")

doc.add_page_break()
add_heading(doc, "数据从哪里来以及怎样防止事后作弊", 1)
add_body(doc, "主要数据来自本地运行的 LSEG 数据环境，包括历史指数成分、公司财报事件、公告前盈利预测、分析师与券商覆盖、价格、收益、成交量和买卖价差。所有原始批次都保留请求身份、文件校验值和运行时间，后续清洗不会覆盖原文件。")

add_heading(doc, "时间切分", 2)
add_table(doc,
          ["区间", "用途", "当前规则"],
          [
              ("2015 至 2020", "训练", "学习模型和选择超参数时使用滚动的年度内部验证"),
              ("2021 至 2022", "验证", "比较模型、检查稳定性和决定最终策略规格"),
              ("2023 至 2026 年 6 月", "最终测试", "保持密封；当前不生成预测，也不查看目标和绩效"),
          ], widths=[Cm(4.2), Cm(4.0), Cm(7.2)])
add_body(doc, "如果一个五日收益窗口跨过训练、验证或测试边界，该事件会被标记为边界隔离，不能进入对应的监督学习样本。这样可以防止训练样本偷偷使用下一阶段的价格。")

add_heading(doc, "不会静默删数据或填数据", 2)
add_body(doc, "每一步都保留状态和原因。例如，公司在公告日不是指数成分、公告前没有足够新的预测快照、价格缺失、入场日不能交易、共同券商不足、五日目标不完整，都会留下明确标记。原始行和主表不会因为“不好看”而消失。")
add_bullet(doc, "缺失值不会自动填成 0；是否可用于模型由单独的 eligibility 标志决定。")
add_bullet(doc, "退市公司和历史指数成分按当时的成员资格处理，避免只保留今天还活着的公司。")
add_bullet(doc, "拆股等公司行动经过单独审计；价格和盈利口径不允许用简单的前后差异硬凑。")
add_bullet(doc, "训练阶段需要的中位数填补只能在每个训练折内部估计，再应用到后面的验证折。")
add_bullet(doc, "所有被替代的运行仍保留，并标记为 superseded，便于追查错误是何时发现和修正的。")

add_heading(doc, "目前的数据规模", 2)
add_table(doc,
          ["数据层", "规模", "它代表什么"],
          [
              ("财报事件主表", "28,995 个事件", "公司财报与公告前一致预期的历史事件集合"),
              ("旧的价格残差网络", "107,532 条公司关系边", "每个事件连接到历史价格行为最相关的候选公司"),
              ("旧网络模型样本", "45 个严格历史特征", "价格、波动、流动性、市场状态和事件特征"),
              ("分析师覆盖原始数据", "46 个季度快照，500,405 行", "覆盖与推荐状态的历史截面"),
              ("新的共同分析师网络", "22,227 个有边事件，111,128 条边", "每个事件最多五个、且公告日成员资格有效的候选公司"),
          ], widths=[Cm(4.3), Cm(4.0), Cm(7.1)])

doc.add_page_break()
add_heading(doc, "目前得到了什么结果", 1)
add_body(doc, "这里有两类结果。第一类是旧网络上的正式训练与验证基线；第二类是新网络上的原始信号快速检查。两者都只使用训练和验证区间，不包含最终测试结果，也都还不是基金回测。")

add_heading(doc, "旧的历史相关性网络", 2)
add_table(doc,
          ["验证模型", "总体排序相关", "同事件内平均排序相关", "加权 R²"],
          [
              ("只看接收公司的 Ridge", "0.0960", "0.0159", "0.0012"),
              ("加入旧网络的 Ridge", "0.0743", "0.0172", "-0.0007"),
              ("加入旧网络的树模型", "0.0209", "-0.0017", "-0.0136"),
          ], widths=[Cm(5.0), Cm(3.4), Cm(4.2), Cm(2.8)])
add_body(doc, "“总体排序相关”把所有样本混在一起，容易受到年份、市场状态或事件整体强弱影响；“同事件内平均排序相关”更接近基金真正需要的能力，因为每次交易都要在同一事件的候选股票中排序。旧网络在这个更关键的指标上只有约 0.016 至 0.017，而且加权 R² 接近 0 或为负。结论是旧网络没有显示出令人信服的额外预测价值。")

add_heading(doc, "新的共同分析师网络", 2)
add_table(doc,
          ["未经模型训练的简单信号", "训练期事件内相关", "验证期事件内相关"],
          [
              ("财报意外 × 覆盖重合度", "0.0046", "0.0101"),
              ("财报意外 × 推荐加权重合度", "0.0052", "0.0118"),
              ("财报意外 × 可识别券商重合度", "0.0040", "-0.0007"),
          ], widths=[Cm(7.3), Cm(4.0), Cm(4.1)])
add_body(doc, "推荐加权信号在训练期和验证期方向一致，验证期约为 0.0118，但幅度很小。跨行业子样本的数字更高，训练期约 0.0136、验证期约 0.0209；不过现有行业标签不是严格的历史时点标签，因此这个结果只能作为诊断，不能拿来筛选交易或挑选模型。")
add_body(doc, "这些数字说明新网络值得进入正式模型比较，却不能说明已经发现可交易的 alpha。快速检查没有控制交易成本、风险暴露和换手，也没有让机器学习学习非线性关系。", bold_lead="可以说什么：")

add_heading(doc, "怎样理解 0.01 左右的排序相关", 2)
add_body(doc, "排序相关等于 1，表示模型每次都把未来表现最好的候选放在最前；等于 0，表示排序和随机差不多；小于 0，则方向相反。量化策略中很小但稳定的相关性有时仍可组合成有用策略，但前提是样本足够、结果跨时期稳定、成本低、风险受控。我们现在只有“值得继续检验”的证据，还没有跨过这些门槛。")

doc.add_page_break()
add_heading(doc, "Claude 版本发现了什么问题以及怎样修正", 1)
add_body(doc, "复核 Claude 新增的共同分析师流程时，我们发现两个会影响研究可信度的问题。旧结果没有被删除，而是保留并明确标记为不可用于决策的版本。")
add_number(doc, "第一版关系网络按覆盖快照检查成员资格，却没有再次确认源公司和接收公司在财报公告当天仍属于股票池。按正确的历史成员区间复核，源公司无效边有 666 条；接收公司在公告日无效的边有 193 条。若只看有入场日期的行，入场日无效边有 194 条。第二版已把公告日成员资格写入硬规则。")
add_number(doc, "第一版快速检查先计算了测试期目标，再在汇总前过滤测试行。虽然没有保存测试指标，这仍违反了密封测试的程序。第二版在连接事件和收益之前就截断到 2022 年末，测试期目标从未生成。")
add_number(doc, "第一版图诊断还使用了测试期的结构信息，并且部分摘要描述超过实际保存内容。该诊断已被判定为 superseded，不再用于模型选择。")

add_heading(doc, "修正后的独立检查", 2)
add_table(doc,
          ["检查对象", "通过数量", "重点检查"],
          [
              ("分析师覆盖原始数据", "18 / 18", "每批请求股票、字段、日期、文件校验值和激活时间"),
              ("共同分析师网络第二版", "21 / 21", "严格事前快照、公告日成员资格、边的数量与相似度重算"),
              ("新信号快速检查第二版", "15 / 15", "测试期截断、边界隔离、目标与排序相关独立重算"),
              ("旧网络模型就绪数据", "31 / 31", "45 个特征、时间边界、交易资格、缺失情况和输入输出校验值"),
          ], widths=[Cm(5.0), Cm(3.0), Cm(7.4)])
add_body(doc, "通过检查表示数据和计算与预先写明的规则一致，不表示投资策略一定有效。研究流程可信和策略有收益是两个不同问题。")

add_heading(doc, "目前哪些东西已经准备好", 1)
add_bullet(doc, "历史股票池、成员起止区间和退市样本处理已经建立。")
add_bullet(doc, "财报事件、公告前预测快照、正式入场会话和五日相对收益定义已经建立。")
add_bullet(doc, "价格、收益、成交量、价差与公司行动检查已有可复现运行。")
add_bullet(doc, "旧价格相关网络、旧模型就绪数据和基线模型已完成。")
add_bullet(doc, "共同分析师覆盖已拉取，第二版关系网络和无测试泄漏的快速检查已完成。")

add_heading(doc, "哪些东西还没有准备好", 1)
add_bullet(doc, "共同分析师网络尚未生成完整的模型就绪数据层。")
add_bullet(doc, "新网络尚未完成接收公司基线、线性模型、树模型和两阶段模型的公平比较。")
add_bullet(doc, "尚未把预测变成每日可交易的持仓，也没有加入仓位、行业、单股、流动性和市场暴露限制。")
add_bullet(doc, "尚未计算换手、佣金、买卖价差、滑点、容量、Sharpe、最大回撤和回撤恢复时间。")
add_bullet(doc, "2023 年以后的最终测试仍然密封，因此目前没有最终样本外表现。")

doc.add_page_break()
add_heading(doc, "接下来怎么把研究变成一只可以被评价的基金", 1)
add_number(doc, "为第二版共同分析师网络建立新的模型就绪数据层。所有测试期特征可以按历史信息计算，但测试期目标保持空白。")
add_number(doc, "在相同训练折和验证期下比较三个层次：只看接收公司、加入共同分析师关系的线性模型、能够学习交互关系的树模型。再增加一个两阶段模型，把事件整体方向和事件内股票排序分开学习。")
add_number(doc, "先用验证期决定唯一的模型、持有期、持仓数量和风险规则。不要因为某个小子样本表现好就反复改条件。")
add_number(doc, "把预测转成组合：每个事件做多排名靠前的股票，必要时做空或用 SPY 对冲，同时限制单股、行业、事件和总市场暴露。")
add_number(doc, "加入可执行性：使用公告后可交易价格、买卖价差、滑点和换手成本，并排除无法按规则成交的行。")
add_number(doc, "冻结全部规则后，只打开一次 2023 年至 2026 年 6 月的测试集，报告收益、Sharpe、回撤、换手、成本敏感性和分时期稳定性。")

add_heading(doc, "投资者最可能追问的风险", 1)
add_table(doc,
          ["风险", "现在怎样控制", "仍需完成"],
          [
              ("数据幸存者偏差", "按历史成员区间保留退市和已移出指数的公司", "继续抽样核对企业行动和标识变化"),
              ("信息穿越", "快照严格早于公告，特征截止入场前，跨边界样本隔离", "新模型就绪层需再次独立复核"),
              ("过拟合", "训练期内滚动选参，2021 至 2022 只做验证，最终测试密封", "冻结少量候选模型与组合规则"),
              ("关系不代表因果", "比较两种关系图，并要求训练与验证方向一致", "做消融、置换和稳定性检验"),
              ("交易成本吞噬信号", "已保留成交量和价差数据", "正式回测加入成本、换手和容量"),
              ("行业或市场暴露伪装成 alpha", "目标扣除 SPY；当前行业结果只作诊断", "组合层加入暴露约束和风险归因"),
          ], widths=[Cm(4.0), Cm(6.2), Cm(5.2)])

add_heading(doc, "当前项目判断", 1)
add_body(doc, "这不是一只已经完成回测的基金，也不是一个只靠复杂模型包装的故事。现阶段最有价值的成果，是我们已经把问题缩小到一个可检验的机制：财报信息是否沿共同分析师覆盖网络传播，并能否帮助我们在同一事件的候选股票中排序。")
add_body(doc, "最初的历史相关性网络可以降为对照组。共同分析师网络应成为主线，但必须通过完整模型、成本后组合和密封测试三个关口。若新网络最终仍只有约 0.01 的不稳定事件内相关，项目应诚实地报告“机制直观但经济价值不足”；若正式模型在验证与测试中稳定提升，并在成本后保留风险调整收益，才可以把它包装成基金的投资边缘。")

add_heading(doc, "术语小词典", 1)
add_table(doc,
          ["术语", "白话解释"],
          [
              ("财报意外", "公司实际公布的盈利与公告前市场一致预期之间的差异，并按预测分歧做标准化。"),
              ("关系边", "源公司和一只候选接收公司之间的一条连接。"),
              ("Jaccard 重合度", "两家公司共同覆盖券商数量，占两家公司全部覆盖券商并集的比例。"),
              ("事件内排序相关", "对同一次财报事件的候选股票，预测排名和未来真实排名有多一致。"),
              ("Ridge", "带有收缩约束的线性回归，适合做稳定、容易解释的基线。"),
              ("密封测试集", "模型和规则完全确定前不允许查看结果的最后一段历史数据。"),
              ("superseded", "旧运行仍保留用于追溯，但已被更正版本替代，不能再用于决策。"),
          ], widths=[Cm(4.0), Cm(11.4)])

# Apply Chinese font and paragraph controls to all runs after content creation.
for p in doc.paragraphs:
    for run in p.runs:
        set_east_asia(run)
    if p.style.name.startswith("Heading"):
        p.paragraph_format.keep_with_next = True

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    set_east_asia(run)
                    run.font.size = Pt(8.5)

doc.core_properties.title = "从一家公司财报推断其他股票的短期反应"
doc.core_properties.subject = "BMF5360 机器学习投资项目当前策略与阶段结果"
doc.core_properties.author = "BMF5360 Project Team"
doc.core_properties.keywords = "Machine Learning, Investments, Earnings, Analyst Network, LSEG"
doc.save(OUT)
print(OUT)
